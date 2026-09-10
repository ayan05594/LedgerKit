"""Convert the maintained LedgerKit reward workbook into deterministic JSON.

Usage:
  python scripts/import-card-research.py path/to/workbook.xlsx \
    src/data/credit-card-research.json

The generated file is consumed by the application. Blank cells deliberately
remain null: the workbook treats a blank as unknown, never as zero or no limit.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import openpyxl


EXPECTED_HEADERS = [
    "Issuer",
    "Status",
    "Card",
    "Base reward",
    "Accelerated rewards",
    "Reward-point value",
    "Monthly/quarterly/annual caps",
    "Minimum transaction",
    "Excluded categories/MCCs",
    "Excluded merchants",
    "Joining fee",
    "Annual fee",
    "Fee-waiver condition",
    "Forex markup",
    "Lounge benefits",
    "Milestone benefits",
    "Fuel surcharge waiver",
    "Other benefits",
    "Official source URL",
    "Effective date",
    "Reward source URL",
    "Reward source type",
    "Reward research date",
    "Reward research note",
]

JSON_KEYS = [
    "issuer",
    "status",
    "name",
    "baseReward",
    "acceleratedRewards",
    "rewardPointValue",
    "caps",
    "minimumTransaction",
    "excludedCategories",
    "excludedMerchants",
    "joiningFee",
    "annualFee",
    "feeWaiverCondition",
    "forexMarkup",
    "loungeBenefits",
    "milestoneBenefits",
    "fuelSurchargeWaiver",
    "otherBenefits",
    "officialSourceUrl",
    "effectiveDate",
    "rewardSourceUrl",
    "rewardSourceType",
    "rewardResearchDate",
    "rewardResearchNote",
]


def clean(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip().replace("\r\n", "\n").replace("\r", "\n")
    return text or None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook_bytes = args.workbook.read_bytes()
    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    if "Catalogue" not in workbook.sheetnames:
        raise SystemExit("Workbook is missing the Catalogue sheet")

    sheet = workbook["Catalogue"]
    values = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(values)]
    if headers != EXPECTED_HEADERS:
        raise SystemExit(f"Unexpected Catalogue headers: {headers!r}")

    cards: list[dict[str, str | None]] = []
    identities: set[tuple[str, str]] = set()
    for row_number, values_row in enumerate(values, start=2):
        values_list = [clean(value) for value in values_row]
        if not any(values_list):
            continue
        record = dict(zip(JSON_KEYS, values_list, strict=True))
        issuer = record["issuer"]
        name = record["name"]
        if not issuer or not name:
            raise SystemExit(f"Row {row_number} is missing issuer or card name")
        identity = (issuer, name)
        if identity in identities:
            raise SystemExit(f"Duplicate card at row {row_number}: {identity!r}")
        identities.add(identity)
        cards.append(record)

    output = {
        "version": "2026-09-09-second-pass",
        "sourceWorkbook": args.workbook.name,
        "sourceSha256": hashlib.sha256(workbook_bytes).hexdigest(),
        "sheet": "Catalogue",
        "cardCount": len(cards),
        "cards": cards,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(output, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
        newline="\n",
    )

    print(f"Wrote {len(cards)} cards to {args.output}")


if __name__ == "__main__":
    main()
